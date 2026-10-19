#!/usr/bin/env python3
"""Build never-roam-alone-packing-list.pdf from packing-list.xlsx (sheet Items, sheet Settings).
Put a URL in the Affiliate URL column to turn a card into an affiliate card
(a rust button, clickable in the PDF). Button label and Note are optional; items stay in section order.
Run: python3 packing-list/build.py"""
import json, html, os, subprocess, tempfile
D=os.path.dirname(os.path.abspath(__file__)); R=os.path.dirname(D)
F="file://"+R.replace(" ","%20")+"/fonts/"
import openpyxl
wb=openpyxl.load_workbook(D+"/packing-list.xlsx",data_only=True)
d={r[0]:r[1] for r in wb["Settings"].iter_rows(values_only=True) if r[0]}
d["sections"]=[]
for sec,name,note,btn,url in (r[:5] for r in wb["Items"].iter_rows(min_row=2,values_only=True)):
    if not name: continue
    if not d["sections"] or d["sections"][-1]["title"]!=sec: d["sections"].append({"title":sec,"items":[]})
    d["sections"][-1]["items"].append({"name":name,"note":note or "","button":btn or "Our pick","url":(url or "").strip()})
e=html.escape
def card(i):
    if isinstance(i,str): i={"name":i}
    note=f"<div class=n>{e(i['note'])}</div>" if i.get("note") else ""
    btn=f"<a class=btn href=\"{e(i['url'])}\">{e(i.get('button') or 'Our pick')} &rarr;</a>" if i.get("url") else ""
    return f"<div class=c><span class=b></span><div class=t><div class=m>{e(i['name'])}</div>{note}{btn}</div></div>"
body="".join(f"<section><h2>{e(s['title'])}</h2><div class=g>{''.join(card(i) for i in s['items'])}</div></section>" for s in d["sections"])
css=f"""
@font-face{{font-family:Fraunces;font-style:italic;font-weight:600;src:url({F}fraunces-italic-600-latin.woff2)}}
@font-face{{font-family:Fraunces;font-weight:700;src:url({F}fraunces-normal-700-latin.woff2)}}
@font-face{{font-family:'Work Sans';font-weight:400;src:url({F}work-sans-normal-400-latin.woff2)}}
@font-face{{font-family:'Work Sans';font-weight:600;src:url({F}work-sans-normal-600-latin.woff2)}}
@page{{size:Letter;margin:0.5in 0.55in}}
*{{box-sizing:border-box;margin:0;padding:0;border-radius:0}}
body{{font-family:'Work Sans',sans-serif;color:#2b2a24;font-size:9.5pt;line-height:1.3}}
header{{border-bottom:3px solid #5c6933;padding-bottom:10px;margin-bottom:12px}}
.k{{font-size:8pt;letter-spacing:.14em;text-transform:uppercase;color:#b5492c;font-weight:600}}
h1{{font-family:Fraunces,serif;font-style:italic;font-weight:600;font-size:25pt;color:#5c6933;line-height:1.1;margin:3px 0 4px}}
.sub{{color:#555}}
section{{margin-bottom:14px}}
h2{{font-family:Fraunces,serif;font-weight:700;font-size:13pt;color:#5c6933;border-bottom:1px solid #cfc7b0;padding-bottom:3px;margin-bottom:7px;break-after:avoid}}
.g{{display:grid;grid-template-columns:1fr 1fr;gap:7px}}
.c{{display:flex;gap:8px;background:#f5efe0;border-left:4px solid #5c6933;padding:7px 10px;break-inside:avoid}}
.b{{flex:none;width:10px;height:10px;border:1.4px solid #b5492c;margin-top:3px;background:#fff}}
.m{{font-weight:600}} .n{{color:#555;font-size:8.7pt;margin-top:1px}}
.btn{{display:inline-block;margin-top:5px;background:#b5492c;color:#fff;text-decoration:none;font-weight:600;font-size:8.3pt;padding:3px 9px}}
footer{{border-top:1px solid #cfc7b0;padding-top:8px;font-size:8.3pt;color:#555}}
footer b{{color:#5c6933}}"""
doc=f"<!doctype html><meta charset=utf-8><title>{e(d['title'])}</title><style>{css}</style><header><div class=k>Never Roam Alone</div><h1>{e(d['title'])}</h1><div class=sub>{e(d['subtitle'])}</div></header>{body}<footer><div>{e(d['disclosure'])}</div><div>Plug types, voltage and weather for every city: <b>neverroamalone.com</b></div></footer>"
out=R+"/never-roam-alone-packing-list.pdf"
with tempfile.TemporaryDirectory() as t:
    open(t+"/p.html","w").write(doc)
    subprocess.run(["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome","--headless","--disable-gpu","--allow-file-access-from-files","--no-pdf-header-footer","--print-to-pdf="+out,"file://"+t+"/p.html"],capture_output=True,check=True)
print(out)
